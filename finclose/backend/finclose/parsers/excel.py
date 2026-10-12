"""Excel(xlsx/xlsm) 판독. 수식과 저장된 계산값을 구분해 읽는다. 매크로는 실행하지 않는다."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .common import ParsedFile, ParseError, build_table, detect_base_date, detect_unit

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MAX_UNCOMPRESSED = 600 * 1024 * 1024


def _precheck(path: Path) -> None:
    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        raise ParseError("비밀번호로 보호(암호화)된 Excel이거나 구형 .xls 형식입니다. 암호를 해제한 .xlsx로 저장 후 업로드하세요.")
    if not head.startswith(b"PK"):
        raise ParseError("Excel(xlsx) 형식이 아니거나 손상된 파일입니다.")
    try:
        with zipfile.ZipFile(path) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_UNCOMPRESSED:
                raise ParseError("압축 해제 크기가 비정상적으로 큽니다(ZIP 폭탄 의심).")
            if "[Content_Types].xml" not in z.namelist():
                raise ParseError("Excel 구조가 아닌 ZIP 파일입니다.")
    except zipfile.BadZipFile as e:
        raise ParseError("손상된 Excel 파일입니다(ZIP 구조 오류).") from e


def parse_excel(path: Path, filename: str, max_rows: int) -> ParsedFile:
    _precheck(path)
    try:
        wb_f = load_workbook(path, data_only=False)
        wb_v = load_workbook(path, data_only=True)
    except ParseError:
        raise
    except Exception as e:  # openpyxl은 다양한 예외를 던진다
        raise ParseError(f"Excel 파일을 열 수 없습니다: {type(e).__name__}") from e

    tables = []
    warnings: list[str] = []
    sheets_meta: list[dict[str, Any]] = []
    all_texts: list[str] = []
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        formula_cells = 0
        formula_no_cache: list[str] = []
        grid: list[list[Any]] = []
        positions: list[str] = []
        num_formats: dict[str, str] = {}
        for r_f, r_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            row_vals = []
            for c_f, c_v in zip(r_f, r_v):
                if c_f.data_type == "f" or (isinstance(c_f.value, str) and c_f.value.startswith("=")):
                    formula_cells += 1
                    if c_v.value is None:
                        formula_no_cache.append(c_f.coordinate)
                if c_f.number_format and c_f.number_format != "General" and c_f.column_letter not in num_formats:
                    num_formats[c_f.column_letter] = c_f.number_format
                row_vals.append(c_v.value)
            grid.append(row_vals)
            positions.append(f"{ws_f.title}!R{r_f[0].row}" if r_f else f"{ws_f.title}!R?")
        hidden_rows = [i for i, d in ws_f.row_dimensions.items() if d.hidden]
        hidden_cols = [k for k, d in ws_f.column_dimensions.items() if d.hidden]
        merged = [str(m) for m in ws_f.merged_cells.ranges]
        smeta = {
            "name": ws_f.title,
            "state": ws_f.sheet_state,
            "max_row": ws_f.max_row,
            "max_column": ws_f.max_column,
            "dimension": ws_f.dimensions,
            "formula_cells": formula_cells,
            "formula_cells_without_cached_value": len(formula_no_cache),
            "formula_no_cache_examples": formula_no_cache[:10],
            "hidden_rows": hidden_rows[:50],
            "hidden_columns": hidden_cols[:50],
            "merged_ranges": merged[:50],
            "number_formats": num_formats,
        }
        if ws_f.sheet_state != "visible":
            warnings.append(f"숨김 시트 '{ws_f.title}'({ws_f.sheet_state})가 있습니다.")
        if formula_no_cache:
            warnings.append(
                f"시트 '{ws_f.title}': 수식 {len(formula_no_cache)}개에 저장된 계산값이 없습니다"
                f"(예: {', '.join(formula_no_cache[:3])}). Excel에서 열어 저장 후 다시 업로드하세요.")
        if hidden_rows or hidden_cols:
            warnings.append(f"시트 '{ws_f.title}'에 숨김 행/열이 있습니다(데이터에는 포함됨).")
        if merged:
            warnings.append(f"시트 '{ws_f.title}'에 병합 셀 {len(merged)}개가 있습니다.")
        texts = [" ".join(str(v) for v in r if v is not None) for r in grid[:15]]
        all_texts.extend(texts)
        if any(any(v is not None for v in r) for r in grid):
            table, tw = build_table(ws_f.title, grid, positions, max_rows, {"sheet_state": ws_f.sheet_state})
            warnings.extend(tw)
            if table:
                table.meta["unit"] = detect_unit(table.meta.get("pre_header_text", []) + table.columns)
                table.meta["formula_cells_without_cached_value"] = len(formula_no_cache)
                smeta["header_row"] = table.header_row
                smeta["data_rows"] = len(table.rows)
                tables.append(table)
        else:
            smeta["empty"] = True
        sheets_meta.append(smeta)
    meta = {
        "sheet_count": len(wb_f.worksheets),
        "sheets": sheets_meta,
        "base_date": detect_base_date(all_texts, filename),
        "unit": detect_unit(all_texts),
    }
    return ParsedFile("excel", tables, meta, warnings)


def column_letter(i: int) -> str:
    return get_column_letter(i + 1)
