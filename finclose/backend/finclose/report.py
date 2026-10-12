"""결과 Excel·HTML 검토 보고서·업로드 파일 생성.

Excel에는 계산된 값만 기록한다(수식 없음) — 수식 캐시 문제를 만들지 않기 위함.
"""
from __future__ import annotations

import csv
import html
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .normalize import to_decimal

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILL = {
    "PASS": "D9F2E3", "WARNING": "FFF2CC", "FAIL": "F8D7DA", "REVIEW_REQUIRED": "E3D9F7",
    "NOT_APPLICABLE": "EDEDED", "NOT_RUN": "FCE4D6",
}
STATUS_KO = {"PASS": "통과", "WARNING": "경고", "FAIL": "실패", "REVIEW_REQUIRED": "검토필요",
             "NOT_APPLICABLE": "해당없음", "NOT_RUN": "미수행"}


def _cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (Decimal, int, float, str, bool)):
        return v
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (list, tuple)):
        return ", ".join(str(_cell(x)) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={_cell(x)}" for k, x in v.items())
    return str(v)


def _write_table(ws, columns: list[str], rows: list[dict[str, Any]], start_row: int = 1,
                 labels: dict[str, str] | None = None) -> int:
    labels = labels or {}
    for j, c in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=labels.get(c, c))
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, r in enumerate(rows, start=start_row + 1):
        for j, c in enumerate(columns, start=1):
            v = _cell(r.get(c))
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, Decimal):
                cell.number_format = "#,##0.########"
            if c == "status" and v in STATUS_FILL:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL[v])
    for j, c in enumerate(columns, start=1):
        width = max([len(str(labels.get(c, c)))] + [len(str(_cell(r.get(c)) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, width * 1.2 + 2))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(rows) + 1


def _cols(rows: list[dict[str, Any]], preferred: list[str] | None = None) -> list[str]:
    cols: list[str] = list(preferred or [])
    for r in rows:
        for k in r:
            if k not in cols and not (k.startswith("_") and k not in ("_src",)):
                cols.append(k)
    return cols


def flatten_details(checks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for c in checks:
        for d in c.get("details", []):
            out.append({"check_id": c["check_id"], "check_name": c["name"], "status": c["status"],
                        "severity": c["severity"], **{k: v for k, v in d.items()}})
    return out


def build_workbook(path: Path, run: dict[str, Any], result: dict[str, Any], gate: dict[str, Any],
                   approvals: list[dict[str, Any]], audit_rows: list[dict[str, Any]], final: bool) -> None:
    tpl_sheets = result.get("output_sheets") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    banner = ("확정(FINALIZED) — 업로드 파일 생성 완료" if final else
              ("검증 차단 사유 있음 — 업로드 금지" if not gate["data_ok"] else "검토/승인 전 — 업로드 금지"))
    ws["A1"] = "FinClose AI 결산 검토 보고서"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = banner
    ws["A2"].font = Font(bold=True, color="C00000" if not final else "006100")
    s = result.get("summary", {})
    info = [
        ("실행 ID", run["id"]), ("업무", f"{run['template_name']} ({run['template_key']} v{run['template_version']})"),
        ("결산 기준일", run["base_date"]), ("실행 상태", run["status"]), ("실행자", run["created_by"]),
        ("실행 시각", run["created_at"]), ("결과 해시", run.get("result_hash")),
        ("원천 건수", s.get("source_count")), ("처리 건수", s.get("processed_count")),
        ("예외 건수", s.get("exception_count")), ("미매칭 건수", s.get("unmatched_count")),
    ]
    for c, t in (s.get("totals_by_currency") or {}).items():
        info += [(f"[{c}] 원천 장부금액 합계", t.get("book_amount_source")), (f"[{c}] 평가금액 합계", t.get("eval_amount")),
                 (f"[{c}] 평가손익 합계", t.get("valuation_gain"))]
    for k in ("total_debit", "total_credit", "voucher_count", "matched", "different", "left_only", "right_only"):
        if k in s:
            info.append((k, s[k]))
    by = {}
    for c in result["checks"]:
        by[c["status"]] = by.get(c["status"], 0) + 1
    info.append(("검증 결과", ", ".join(f"{STATUS_KO[k]} {v}" for k, v in by.items())))
    info.append(("차단 사유", "; ".join(b["reason"] for b in gate["blockers"]) or "없음"))
    info.append(("검토 확인 대기", "; ".join(p["reason"] for p in gate["pending_reviews"]) or "없음"))
    r = 4
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=_cell(v))
        if isinstance(v, Decimal):
            cell.number_format = "#,##0.########"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="사용된 원천 파일").font = Font(bold=True)
    r += 1
    r = _write_table(ws, ["role", "label", "file", "sha256", "table", "rows"],
                     [{"role": k, **v} for k, v in result["inputs"].items()], r) + 1
    ws.cell(row=r, column=1, value="계산 규칙 버전").font = Font(bold=True)
    r = _write_table(ws, ["rule_id", "version", "approval_status"], result.get("rule_versions", []), r + 1) + 1
    if result.get("confirm_items"):
        ws.cell(row=r, column=1, value="실제 적용 전 확인 필요 사항").font = Font(bold=True)
        for i, item in enumerate(result["confirm_items"], start=r + 1):
            ws.cell(row=i, column=1, value=f"• {item}")
        r = i + 2
    ws.cell(row=r, column=1, value="승인 이력").font = Font(bold=True)
    _write_table(ws, ["at", "action", "check_id", "actor", "role", "comment"], approvals, r + 1)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    ws.freeze_panes = None

    def want(name: str) -> bool:
        return not tpl_sheets or name in tpl_sheets

    if want("Source_Data"):
        wsx = wb.create_sheet("Source_Data")
        row = 1
        for role, recs in result.get("source_data", {}).items():
            wsx.cell(row=row, column=1, value=f"[{role}] {result['inputs'][role].get('label')}").font = Font(bold=True)
            row = _write_table(wsx, _cols(recs, ["_src"]), recs, row + 1) + 2
    if want("Processed_Data"):
        rows = result.get("processed", [])
        _write_table(wb.create_sheet("Processed_Data"), _cols(rows, ["_src"]) + (["_trace"] if rows and "_trace" in rows[0] else []), rows)
    if want("Reconciliation") and result.get("reconciliation"):
        rows = result["reconciliation"]
        _write_table(wb.create_sheet("Reconciliation"), _cols(rows), rows)
    if want("Validation"):
        _write_table(wb.create_sheet("Validation"),
                     ["check_id", "category", "name", "status", "severity", "mandatory", "message", "method"],
                     [{**c, "status": c["status"]} for c in result["checks"]])
    if want("Variance_Analysis"):
        wsv = wb.create_sheet("Variance_Analysis")
        row = 1
        for code, v in (result.get("variance", {}).get("periods") or {}).items():
            wsv.cell(row=row, column=1, value=f"{v['label']} (비교 기준일 {v['base_date']}, 출처: {v.get('source') or '-'})").font = Font(bold=True)
            if v.get("result"):
                t = v["result"]["total"]
                wsv.cell(row=row + 1, column=1, value=f"합계: 당기 {t['current']:,} / 비교 {t['base']:,} / 증감 {t['delta']:,} / {t['rate_status']}"
                         + (f" {t['rate']:.2%}" if t.get("rate") is not None else ""))
                row = _write_table(wsv, ["key", "label", "current", "base", "delta", "rate_status", "rate", "current_qty",
                                         "base_qty", "share_change", "flags", "cause_candidates"], v["result"]["rows"], row + 2) + 2
            else:
                wsv.cell(row=row + 1, column=1, value=f"미수행: {v.get('not_run_reason')}")
                row += 3
        trend = result.get("variance", {}).get("trend") or []
        if trend:
            wsv.cell(row=row, column=1, value="최근 추이").font = Font(bold=True)
            _write_table(wsv, ["base_date", "total", "count", "run_id"], trend, row + 1)
    if want("Exceptions"):
        wse = wb.create_sheet("Exceptions")
        row = 1
        exc = result.get("exceptions", [])
        wse.cell(row=row, column=1, value=f"처리 제외 행 ({len(exc)}건)").font = Font(bold=True)
        row = _write_table(wse, _cols(exc, ["_src", "reason"]), exc, row + 1) + 2
        det = flatten_details([c for c in result["checks"] if c["status"] not in ("PASS", "NOT_APPLICABLE")])
        wse.cell(row=row, column=1, value=f"검증 상세 ({len(det)}건)").font = Font(bold=True)
        _write_table(wse, _cols(det, ["check_id", "check_name", "status", "severity", "role", "file", "source_row",
                                      "column", "source_value", "computed", "compared", "difference", "formula", "reason"]),
                     det, row + 1)
    if result.get("journal") and want("Journal_Draft"):
        wsj = wb.create_sheet("Journal_Draft")
        wsj["A1"] = "분개 초안 — 원장 미반영(DRAFT). 계정은 예시일 수 있음."
        wsj["A1"].font = Font(bold=True, color="C00000")
        _write_table(wsj, _cols(result["journal"]), result["journal"], 2)
    if want("Upload_Data"):
        wsu = wb.create_sheet("Upload_Data")
        wsu["A1"] = "확정된 업로드 대상" if final else "미확정(DRAFT) 미리보기 — 업로드 금지"
        wsu["A1"].font = Font(bold=True, color="006100" if final else "C00000")
        _write_table(wsu, result.get("upload_columns_all") or _cols(result.get("upload_rows", [])), result.get("upload_rows", []), 2)
    if want("Audit_Log"):
        _write_table(wb.create_sheet("Audit_Log"), ["id", "at", "actor", "role", "action", "entity_id", "details_json", "hash"], audit_rows)
    for wsx in wb.worksheets:
        for row in wsx.iter_rows(min_row=1, max_row=1):
            for c in row:
                c.alignment = Alignment(vertical="center")
    tmp = path.with_suffix(".xlsx.part")
    wb.save(tmp)
    tmp.replace(path)


def write_upload_files(dir_: Path, stem: str, columns: list[str], rows: list[dict[str, Any]]) -> dict[str, Path]:
    """표준 형식 업로드 파일(CSV UTF-8 BOM, XLSX). 대상 시스템 규격 어댑터는 미구현."""
    csv_path, xlsx_path = dir_ / f"{stem}.csv", dir_ / f"{stem}.xlsx"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(columns)
        for r in rows:
            w.writerow([_cell(r.get(c)) if r.get(c) is not None else "" for c in columns])
    wb = Workbook()
    ws = wb.active
    ws.title = "Upload"
    ws.append(columns)
    for r in rows:
        ws.append([_cell(r.get(c)) for c in columns])
    wb.save(xlsx_path)
    return {"csv": csv_path, "xlsx": xlsx_path}


def read_back_totals(xlsx_path: Path, csv_path: Path, columns: list[str], amount_cols: list[str]) -> dict[str, Any]:
    """기록한 업로드 파일을 다시 읽어 건수·합계를 산출(T05 검증용 — 기록 경로와 별개)."""
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    out: dict[str, Any] = {"xlsx_count": len(rows) - 1, "xlsx_sums": {}, "csv_sums": {}}
    for c in amount_cols:
        i = header.index(c)
        out["xlsx_sums"][c] = sum((to_decimal(r[i])[0] or Decimal(0) for r in rows[1:]), Decimal(0))
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        cr = list(csv.reader(f))
    out["csv_count"] = len(cr) - 1
    for c in amount_cols:
        i = cr[0].index(c)
        out["csv_sums"][c] = sum((to_decimal(r[i])[0] or Decimal(0) for r in cr[1:]), Decimal(0))
    return out


def build_html(run: dict[str, Any], result: dict[str, Any], gate: dict[str, Any], approvals: list[dict[str, Any]]) -> str:
    e = html.escape
    s = result.get("summary", {})
    rows = "".join(
        f"<tr class='{c['status']}'><td>{e(c['check_id'])}</td><td>{e(c['category'])}</td><td>{e(c['name'])}</td>"
        f"<td>{STATUS_KO[c['status']]}</td><td>{e(c['severity'])}</td><td>{e(c['message'])}</td></tr>"
        for c in result["checks"])
    var = ""
    for v in (result.get("variance", {}).get("periods") or {}).values():
        if v.get("result"):
            t = v["result"]["total"]
            rate = f"{t['rate']:.2%}" if t.get("rate") is not None else t["rate_status"]
            var += (f"<tr><td>{e(v['label'])}</td><td>{e(v['base_date'])}</td><td class=n>{t['current']:,}</td>"
                    f"<td class=n>{t['base']:,}</td><td class=n>{t['delta']:,}</td><td>{e(rate)}</td>"
                    f"<td class=n>{v['result']['anomaly_count']}</td></tr>")
        else:
            var += f"<tr><td>{e(v['label'])}</td><td>{e(v['base_date'])}</td><td colspan=5>미수행: {e(v.get('not_run_reason') or '')}</td></tr>"
    files = "".join(f"<li>{e(k)}: {e(str(v.get('file')))} <code>{e(str(v.get('sha256'))[:16])}…</code> ({v.get('rows')}행)</li>"
                    for k, v in result["inputs"].items())
    totals = "".join(f"<li>{e(c)}: 장부 {t['book_amount_source']:,} / 평가 {t['eval_amount']:,} / 평가손익 {t['valuation_gain']:,} ({t['count']}건)</li>"
                     for c, t in (s.get("totals_by_currency") or {}).items())
    blockers = "".join(f"<li>{e(b['reason'])}</li>" for b in gate["blockers"] + gate["pending_reviews"]) or "<li>없음</li>"
    appr = "".join(f"<li>{e(a['at'])} {e(a['action'])} {e(a.get('check_id') or '')} — {e(a['actor'])}({e(a['role'])}): {e(a['comment'])}</li>"
                   for a in approvals) or "<li>없음</li>"
    confirm = "".join(f"<li>{e(x)}</li>" for x in result.get("confirm_items", []))
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>결산 검토 보고서 {e(run['id'])}</title>
<style>body{{font-family:'Malgun Gothic',sans-serif;margin:24px;color:#1b2430}}table{{border-collapse:collapse;width:100%}}
td,th{{border:1px solid #ccd;padding:4px 6px;font-size:13px}}th{{background:#1f3a5f;color:#fff}}.n{{text-align:right}}
.PASS{{background:#d9f2e3}}.FAIL{{background:#f8d7da}}.WARNING{{background:#fff2cc}}.REVIEW_REQUIRED{{background:#e3d9f7}}
.NOT_RUN{{background:#fce4d6}}.NOT_APPLICABLE{{background:#ededed}}.banner{{padding:8px;font-weight:bold;border:2px solid #c00;color:#c00}}</style></head>
<body><h1>결산 검토 보고서</h1><div class="banner">{'확정 완료' if run['status']=='FINALIZED' else '미확정 — 업로드 금지'}</div>
<p><b>업무</b> {e(run['template_name'])} ({e(run['template_key'])} v{run['template_version']}) · <b>기준일</b> {e(run['base_date'])}
 · <b>상태</b> {e(run['status'])} · <b>실행</b> {e(run['id'])} by {e(run['created_by'])} @ {e(run['created_at'])}</p>
<h2>원천 파일</h2><ul>{files}</ul>
<h2>처리 요약</h2><p>원천 {s.get('source_count')}건 · 처리 {s.get('processed_count')}건 · 예외 {s.get('exception_count')}건 · 미매칭 {s.get('unmatched_count', '-')}건</p><ul>{totals}</ul>
<h2>차단 사유 / 검토 필요</h2><ul>{blockers}</ul>
<h2>검증 결과</h2><table><tr><th>ID</th><th>범주</th><th>검증</th><th>상태</th><th>심각도</th><th>내용</th></tr>{rows}</table>
<h2>기간 비교</h2><table><tr><th>구분</th><th>비교 기준일</th><th>당기</th><th>비교</th><th>증감</th><th>증감률</th><th>이상 후보</th></tr>{var}</table>
<h2>실제 적용 전 확인 사항</h2><ul>{confirm or '<li>-</li>'}</ul>
<h2>승인 이력</h2><ul>{appr}</ul>
<p style="color:#666;font-size:12px">이 보고서의 증감 원인은 추정 후보이며 원천 거래자료로 확인되지 않았습니다. 분개는 초안이며 원장에 반영되지 않았습니다.</p>
</body></html>"""
