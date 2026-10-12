from pathlib import Path

import pytest
from openpyxl import Workbook

from finclose.parsers import ParseError, parse_file
from finclose.profiler import profile_table


def _xlsx(path: Path, build) -> Path:
    wb = Workbook()
    build(wb)
    wb.save(path)
    return path


def test_excel_metadata_title_rows_unit_formula_hidden_merged(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "데이터"
        ws["A1"] = "보유 채권 내역"
        ws.merge_cells("A1:C1")
        ws["A2"] = "기준일: 2026-09-30"
        ws["D2"] = "(단위: 천원)"
        ws.append([])
        ws.append(["종목코드", "수량", "금액"])
        ws.append(["KR9TEST00001", 10, 100])
        ws.append(["KR9TEST00002", 20, "=B6*10"])
        ws.append(["합계", 30, 300])
        ws.row_dimensions[5].hidden = True
        h = wb.create_sheet("숨김")
        h.sheet_state = "hidden"
        h.append(["a", "b"])
    p = _xlsx(tmp_path / "x.xlsx", build)
    pf = parse_file(p, "xlsx", "x.xlsx")
    t = pf.tables[0]
    assert t.columns == ["종목코드", "수량", "금액"]
    assert t.header_row == 4 and len(t.rows) == 2            # 합계 행 제외
    assert t.meta["unit"]["label"] == "천원"
    assert pf.metadata["base_date"]["date"] == "2026-09-30"
    sm = pf.metadata["sheets"][0]
    assert sm["formula_cells"] == 1 and sm["formula_cells_without_cached_value"] == 1
    assert sm["merged_ranges"] == ["A1:C1"] and 5 in sm["hidden_rows"]
    assert any("숨김 시트" in w for w in pf.warnings)
    assert any("저장된 계산값이 없습니다" in w for w in pf.warnings)
    assert t.source_rows[0] == "데이터!R5"


def test_corrupted_xlsx(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"PK\x03\x04garbage-not-a-zip" * 10)
    with pytest.raises(ParseError, match="손상"):
        parse_file(p, "xlsx", "bad.xlsx")


def test_encrypted_or_legacy_xls_detected(tmp_path):
    p = tmp_path / "enc.xlsx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
    with pytest.raises(ParseError, match="비밀번호"):
        parse_file(p, "xlsx", "enc.xlsx")


def test_empty_file(tmp_path):
    p = tmp_path / "e.csv"
    p.write_bytes(b"")
    with pytest.raises(ParseError, match="빈 파일"):
        parse_file(p, "csv", "e.csv")


def test_csv_cp949(tmp_path):
    p = tmp_path / "k.csv"
    p.write_bytes("계정코드,계정명,금액\n1101,현금,\"1,000\"\n".encode("cp949"))
    pf = parse_file(p, "csv", "k.csv")
    assert pf.metadata["encoding"] == "cp949"
    assert pf.tables[0].rows == [["1101", "현금", "1,000"]]


def test_csv_without_header_reports_warning(tmp_path):
    p = tmp_path / "n.csv"
    p.write_text("1,2\n3,4\n")
    pf = parse_file(p, "csv", "n.csv")
    assert pf.tables == [] and any("헤더" in w for w in pf.warnings)


def test_pdf_table_extraction(samples):
    pf = parse_file(samples["pdf"], "pdf", samples["pdf"].name)
    t = pf.tables[0]
    assert t.columns[:2] == ["Code", "Agency"]
    assert len(t.rows) == 22
    assert t.meta["review_status"] == "EXTRACTED_UNVERIFIED"
    assert t.source_rows[0].startswith("p1 ")


def test_pdf_without_tables_and_invalid_pdf(tmp_path):
    from reportlab.pdfgen import canvas
    p = tmp_path / "t.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(100, 750, "No table here")
    c.save()
    pf = parse_file(p, "pdf", "t.pdf")
    assert pf.tables == [] and any("표를 추출하지 못했습니다" in w for w in pf.warnings)
    bad = tmp_path / "b.pdf"
    bad.write_bytes(b"not a pdf at all")
    with pytest.raises(ParseError):
        parse_file(bad, "pdf", "b.pdf")


def test_row_limit(tmp_path):
    p = tmp_path / "big.csv"
    p.write_text("a,b\n" + "x,1\n" * 30)
    with pytest.raises(ParseError, match="제한"):
        parse_file(p, "csv", "big.csv", max_rows=10)


def test_profile_detects_missing_dups_mixed(tmp_path):
    p = tmp_path / "p.csv"
    p.write_text("code,amt\nA,1\nA,2\n,x\nB, 3 \nC,1000000\nD,5\nE,6\nF,7\nG,8\n")
    t = parse_file(p, "csv", "p.csv").tables[0]
    prof = profile_table(t)
    assert prof["columns"]["code"]["missing"] == 1
    assert prof["columns"]["code"]["duplicate_rows"] == 2
    amt = prof["columns"]["amt"]
    assert amt["mixed_types"] and amt["non_numeric_values"] == 1
    assert amt["leading_trailing_spaces"] == 1
    assert amt["outlier_candidates"] >= 1


def test_pdf_borderless_table_fallback(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table
    p = tmp_path / "nb.pdf"
    SimpleDocTemplate(str(p), pagesize=A4).build([Table([["Code", "Price", "CCY"], ["KR9TEST00001", "9999.10", "KRW"],
                                                         ["KR9TEST00002", "10001.00", "KRW"]])])
    pf = parse_file(p, "pdf", "nb.pdf")
    t = pf.tables[0]
    assert t.columns == ["Code", "Price", "CCY"] and len(t.rows) == 2
    assert t.meta["extraction_method"] == ["text_alignment"]
