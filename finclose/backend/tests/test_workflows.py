"""통합 테스트: 업로드 → 분석 → 실행 → 검증 → 증감 → 승인 → 확정 및 실패 시나리오."""
import errno
import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finclose import intake, report
from finclose import templates as treg
from finclose.service import Service, ServiceError

from .conftest import approve_and_finalize, bond_run, check, upload


def test_happy_path_analyze_run_finalize(svc, samples):
    ids = upload(svc, samples["current"].values())
    a = svc.analyze(ids)
    best = a["candidates"][0]
    assert best["template_key"] == "BOND_EXT_PRICE_RECON" and best["is_candidate"]
    assert set(best["assignments"]) == {"holdings", "prices", "prior"}
    plan = a["plan"]
    assert plan["base_date"] == "2026-09-30" and not plan["needs_confirmation"]
    files = {r: e["file_id"] for r, e in plan["roles"].items()}
    run = svc.create_run("BOND_EXT_PRICE_RECON", None, plan["base_date"], files, {}, "이담당", "preparer")
    assert run["status"] == "READY_FOR_REVIEW"
    for cid in ("F01", "F03", "D02", "R01", "T01", "T02", "T03", "T04", "C01", "C02", "C03", "C04", "C05", "A01", "A02", "R03"):
        assert check(run, cid)["status"] == "PASS", cid
    assert check(run, "C06")["status"] == "REVIEW_REQUIRED"        # 예시 산식
    assert check(run, "A03")["status"] == "REVIEW_REQUIRED"        # 예시 계정
    assert run["summary"]["processed_count"] == 22 and run["summary"]["exception_count"] == 0
    # 전월 파일 비교: 신규 2건(23,24번) / 소멸 1건(11번)
    rows = run["variance"]["periods"]["PRIOR_FILE"]["result"]["rows"]
    st = {r["key"].split("|")[0]: r["rate_status"] for r in rows}
    assert st["KR9TEST00023"] == "NEW" and st["KR9TEST00011"] == "REMOVED"
    final = approve_and_finalize(svc, run)
    assert final["status"] == "FINALIZED"
    assert {"final.xlsx", "upload_BOND_EXT_PRICE_RECON_20260930.csv"} <= set(final["downloads"])
    wb = load_workbook(svc.download_path(run["id"], "final.xlsx"))
    assert {"Summary", "Source_Data", "Processed_Data", "Reconciliation", "Validation", "Variance_Analysis",
            "Exceptions", "Journal_Draft", "Upload_Data", "Audit_Log"} <= set(wb.sheetnames)
    # 업로드 파일 합계 = 처리 결과 합계
    up = load_workbook(svc.download_path(run["id"], "upload_BOND_EXT_PRICE_RECON_20260930.xlsx")).active
    hdr = [c.value for c in up[1]]
    i = hdr.index("eval_amount")
    from decimal import Decimal
    total = sum(Decimal(str(r[i])) for r in up.iter_rows(min_row=2, values_only=True))
    expected = sum(Decimal(v["eval_amount"]) for v in run["summary"]["totals_by_currency"].values())
    assert total == expected


def test_reproducible_and_originals_untouched(svc, samples):
    paths = samples["current"]
    before = {r: hashlib.sha256(Path(p).read_bytes()).hexdigest() for r, p in paths.items()}
    r1 = bond_run(svc, paths)
    r2 = bond_run(svc, paths, rerun_reason="재현성 확인")   # 같은 해시라도 업로드 id가 달라 fingerprint 동일
    assert r1["result_hash"] == r2["result_hash"]
    after = {r: hashlib.sha256(Path(p).read_bytes()).hexdigest() for r, p in paths.items()}
    assert before == after
    for f in svc.db.query("SELECT * FROM files"):
        assert intake.file_sha_ok(f)
        assert (Path(f["stored_path"]).stat().st_mode & 0o222) == 0     # 읽기전용


def test_error_dataset_detected_and_blocked(svc, samples):
    run = bond_run(svc, samples["errors"])
    assert run["status"] == "BLOCKED"
    expect = {"D02": "FAIL", "D01": "FAIL", "D04": "FAIL", "D07": "FAIL", "F05": "FAIL", "R01": "FAIL", "C03": "FAIL",
              "D08": "WARNING", "R02": "WARNING"}
    for cid, st in expect.items():
        assert check(run, cid)["status"] == st, cid
    reasons = {e["code"]: e["reason"] for e in run["exceptions"]}
    assert "미매칭" in reasons["KR9TEST99901"] and "음수" in reasons["KR9TEST00006"] and "누락" in reasons["KR9TEST00010"]
    # T01/T02 still hold: 예외도 건수·금액에 포함
    assert check(run, "T01")["status"] == "PASS" and check(run, "T02")["status"] == "PASS"
    d = check(run, "R01")["details"][0]
    assert d["source_row"].startswith("보유채권내역_20260930_오류.xlsx 보유채권!R")
    with pytest.raises(ServiceError) as e:
        svc.finalize(run["id"], "확정 시도", "박승인", "approver")
    assert e.value.status == 409
    assert not list((svc.settings.runs_dir / run["id"]).glob("upload_*"))
    # CRITICAL은 예외승인 불가, HIGH는 가능하며 원래 결과 보존
    with pytest.raises(ServiceError, match="CRITICAL"):
        svc.act(run["id"], "EXCEPTION", "D02", "무시", "박승인", "approver")
    r = svc.act(run["id"], "EXCEPTION", "R01", "다음 달 단가 수령 예정", "박승인", "approver")
    assert check(r, "R01")["status"] == "FAIL"
    assert all(b["check_id"] != "R01" for b in r["gate"]["blockers"])
    assert any(a["action"] == "EXCEPTION" and a["check_id"] == "R01" for a in r["approvals"])


def test_missing_required_file(svc, samples):
    p = samples["current"]
    run = bond_run(svc, {"holdings": p["holdings"], "prices": p["prices"]})
    assert check(run, "F01")["status"] == "FAIL" and check(run, "F01")["severity"] == "CRITICAL"
    assert run["status"] == "BLOCKED"


def test_wrong_base_date_parameter(svc, samples):
    run = bond_run(svc, samples["current"], base="2026-10-31")
    assert check(run, "F05")["status"] == "FAIL"


def test_unit_error_thousand_won(svc, samples):
    p = dict(samples["current"])
    p["holdings"] = samples["thousand"]
    run = bond_run(svc, p)
    assert check(run, "F06")["status"] == "FAIL"
    assert check(run, "C03")["status"] == "FAIL"          # 평가/장부 비율로도 탐지


def test_duplicate_run_rejected(svc, samples):
    paths = samples["current"]
    roles = {r: upload(svc, [p])[0] for r, p in paths.items()}
    svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30", roles, {}, "이담당", "preparer")
    with pytest.raises(ServiceError) as e:
        svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30", roles, {}, "이담당", "preparer")
    assert e.value.payload["code"] == "DUPLICATE_RUN"


def test_double_finalize_blocked(svc, samples):
    r1 = approve_and_finalize(svc, bond_run(svc, samples["current"]))
    r2 = bond_run(svc, samples["current"], rerun_reason="재실행")
    assert any("중복 반영" in b["reason"] for b in r2["gate"]["blockers"])
    assert r1["status"] == "FINALIZED"


def test_segregation_of_duties_and_permissions(svc, samples):
    run = bond_run(svc, samples["current"])
    with pytest.raises(ServiceError, match="직무분리"):
        svc.act(run["id"], "REVIEW", None, "셀프검토", "이담당", "reviewer")
    with pytest.raises(ServiceError, match="REVIEWED"):
        svc.act(run["id"], "APPROVE", None, "바로승인", "박승인", "approver")
    with pytest.raises(ServiceError, match="사유"):
        svc.act(run["id"], "REVIEW", None, "", "김검토", "reviewer")


def test_approve_blocked_when_review_items_unacknowledged(svc, samples):
    run = bond_run(svc, samples["current"])
    svc.act(run["id"], "REVIEW", None, "검토", "김검토", "reviewer")
    with pytest.raises(ServiceError, match="차단"):
        svc.act(run["id"], "APPROVE", None, "승인", "박승인", "approver")


def test_history_comparisons_and_format_change(svc, samples):
    for key in sorted(k for k in samples if k.startswith("history_")):
        hd = key.split("_", 1)[1]
        r = bond_run(svc, samples[key], base=hd)
        assert r["status"] == "READY_FOR_REVIEW", (hd, r["gate"])
        approve_and_finalize(svc, r)
    cur = bond_run(svc, samples["current"])
    per = cur["variance"]["periods"]
    for code in ("PRIOR_FILE", "PREV_MONTH", "PREV_QUARTER", "SAME_PERIOD_LAST_YEAR", "PRIOR_YEAR_END"):
        assert per[code].get("result"), code
    assert len(cur["variance"]["trend"]) == 4
    assert check(cur, "F07")["status"] == "PASS"
    # 수량 2배 종목(5번): 수량 변동 원인 후보
    row = next(r for r in per["PREV_MONTH"]["result"]["rows"] if r["key"].startswith("KR9TEST00005"))
    assert any("보유수량 변동" in c for c in row["cause_candidates"])
    # 형식 변경: 자동 처리 중단 → 사유와 함께 확인 시 진행, F07 REVIEW_REQUIRED
    p = dict(samples["current"])
    p["holdings"] = samples["renamed"]
    roles = {r: upload(svc, [x])[0] for r, x in p.items()}
    with pytest.raises(ServiceError) as e:
        svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30", roles, {}, "이담당", "preparer")
    assert e.value.status == 409 and e.value.payload["code"] == "NEEDS_CONFIRMATION"
    kinds = {d["type"] for d in e.value.payload["drift"]["holdings"]}
    assert {"SHEET_RENAMED", "COLUMN_ADDED"} <= kinds
    run = svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30", roles, {}, "이담당", "preparer",
                         confirm_changes=True, confirm_reason="시스템 개편으로 컬럼명 변경 확인")
    assert check(run, "F07")["status"] == "REVIEW_REQUIRED"


def test_crash_recovery(settings, samples):
    svc = Service(settings)
    svc.db.execute("INSERT INTO runs(id, template_key, template_version, base_date, status, created_at, created_by,"
                   " input_files_json, fingerprint) VALUES ('RX','BOND_EXT_PRICE_RECON',1,'2026-09-30','RUNNING','t','u','{}','fp')")
    (settings.runs_dir / "RX.tmp").mkdir()
    (settings.upload_dir / "zz.xlsx.part").write_bytes(b"x")
    svc.db.close()
    svc2 = Service(settings)
    r = svc2.db.one("SELECT status, error FROM runs WHERE id='RX'")
    assert r["status"] == "FAILED" and "비정상 종료" in r["error"]
    assert not (settings.runs_dir / "RX.tmp").exists() and not (settings.upload_dir / "zz.xlsx.part").exists()


def test_disk_full_on_upload(svc, samples, monkeypatch):
    real = shutil.disk_usage
    monkeypatch.setattr(shutil, "disk_usage", lambda p: real(p)._replace(free=0))
    with open(samples["current"]["holdings"], "rb") as f, pytest.raises(intake.IntakeError) as e:
        intake.save_upload(svc.db, svc.settings, f, "a.xlsx", "u", "preparer")
    assert e.value.status == 507


def test_enospc_during_write_cleans_partial(svc, monkeypatch):
    import io

    class Boom(io.BytesIO):
        def read(self, n=-1):
            raise OSError(errno.ENOSPC, "No space left on device")
    with pytest.raises(intake.IntakeError) as e:
        intake.save_upload(svc.db, svc.settings, Boom(b"x"), "a.csv", "u", "preparer")
    assert e.value.status == 507
    assert not list(svc.settings.upload_dir.glob("*.part"))


@pytest.mark.parametrize("name,data,msg", [
    ("a.exe", b"MZ", "허용되지 않는"), ("old.xls", b"x", ".xls"), ("empty.csv", b"", "빈 파일"),
    ("../../etc/passwd.csv", b"a,b\n1,2\n", None),
])
def test_upload_validation(svc, name, data, msg):
    import io
    if msg:
        with pytest.raises(intake.IntakeError, match=msg):
            intake.save_upload(svc.db, svc.settings, io.BytesIO(data), name, "u", "preparer")
    else:
        rec = intake.save_upload(svc.db, svc.settings, io.BytesIO(data), name, "u", "preparer")
        assert rec["original_name"] == "passwd.csv"
        assert Path(rec["stored_path"]).parent == svc.settings.upload_dir


def test_corrupted_excel_upload_cannot_run(svc, tmp_path, samples):
    import io
    rec = intake.save_upload(svc.db, svc.settings, io.BytesIO(b"PK\x03\x04broken" * 20), "보유채권내역.xlsx", "u", "preparer")
    assert rec["status"] == "ERROR" and "손상" in rec["error"]
    ids = upload(svc, [samples["current"]["prices"], samples["current"]["prior"]])
    with pytest.raises(ServiceError, match="판독 오류"):
        svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30",
                       {"holdings": rec["id"], "prices": ids[0], "prior": ids[1]}, {}, "u", "preparer")


def test_tampered_original_refused(svc, samples):
    ids = upload(svc, samples["current"].values())
    f = svc.db.one("SELECT stored_path FROM files WHERE id=?", (ids[0],))
    p = Path(f["stored_path"])
    p.chmod(0o644)
    p.write_bytes(p.read_bytes() + b"x")
    roles = dict(zip(["holdings", "prices", "prior"], ids))
    with pytest.raises(ServiceError, match="변조"):
        svc.create_run("BOND_EXT_PRICE_RECON", None, "2026-09-30", roles, {}, "u", "preparer")


def test_readback_failure_discards_upload_files(svc, samples, monkeypatch):
    run = bond_run(svc, samples["current"])
    rid = run["id"]
    for p in run["gate"]["pending_reviews"]:
        svc.act(rid, "ACKNOWLEDGE", p["check_id"], "확인", "김검토", "reviewer")
    svc.act(rid, "REVIEW", None, "검토", "김검토", "reviewer")
    svc.act(rid, "APPROVE", None, "승인", "박승인", "approver")
    real = report.write_upload_files

    def corrupt(dir_, stem, cols, rows):
        return real(dir_, stem, cols, rows[:-1])        # 한 행 누락된 파일
    monkeypatch.setattr("finclose.service.write_upload_files", corrupt)
    with pytest.raises(ServiceError, match="T05"):
        svc.finalize(rid, "확정", "박승인", "approver")
    assert not list((svc.settings.runs_dir / rid).glob("upload_*"))
    assert svc.get_run(rid)["status"] == "APPROVED"


def test_download_path_traversal(svc, samples):
    run = bond_run(svc, samples["current"])
    for bad in ("../finclose.db", "result.json", "../../etc/passwd"):
        with pytest.raises(ServiceError):
            svc.download_path(run["id"], bad)


def test_invalid_template_rule_rejected_and_activation_requires_tests(svc):
    body = treg.get(svc.db, "BOND_EXT_PRICE_RECON")["body"]
    bad = {**body, "calculations": [{**body["calculations"][0], "expression": "quantity * __import__"}]}
    with pytest.raises(treg.TemplateRegistryError):
        treg.save_version(svc.db, bad, "이담당", "preparer", "잘못된 규칙")
    wrong = {**body, "calculations": [{**body["calculations"][0],
                                       "test_cases": [{"inputs": {"quantity": "1", "price": "10000"}, "currency": "KRW", "expected": "999"}]}]}
    t = treg.save_version(svc.db, wrong, "이담당", "preparer", "기대값 오류")
    assert t["status"] == "DRAFT" and not t["test_result"]["ok"]
    with pytest.raises(treg.TemplateRegistryError, match="테스트"):
        treg.set_status(svc.db, "BOND_EXT_PRICE_RECON", t["version"], "ACTIVE", "박승인", "approver", "활성화")
    assert treg.get(svc.db, "BOND_EXT_PRICE_RECON")["version"] == 1      # 기존 활성 버전 유지


def test_preparer_cannot_self_approve_rules(svc):
    body = treg.get(svc.db, "BOND_EXT_PRICE_RECON")["body"]
    esc = {**body, "calculations": [{**c, "approval_status": "APPROVED"} for c in body["calculations"]]}
    with pytest.raises(treg.TemplateRegistryError, match="승인자"):
        treg.save_version(svc.db, esc, "이담당", "preparer", "승인처리")
    t = treg.save_version(svc.db, esc, "박승인", "approver", "정책위원회 승인(가정)")
    t = treg.set_status(svc.db, "BOND_EXT_PRICE_RECON", t["version"], "ACTIVE", "박승인", "approver", "활성")
    assert treg.get(svc.db, "BOND_EXT_PRICE_RECON", 1)["status"] == "SUPERSEDED"   # 이전 버전 보존


def test_template_change_versioning_and_clone(svc):
    t = treg.clone(svc.db, "BOND_EXT_PRICE_RECON", 1, "BOND_COPY_TEST", "복제본", "이담당", "preparer")
    assert t["version"] == 1 and t["status"] == "DRAFT"
    with pytest.raises(treg.TemplateRegistryError, match="사유"):
        treg.save_version(svc.db, t["body"], "이담당", "preparer", "")


def test_journal_check_valid_and_unbalanced(svc, samples):
    j = samples["journal"]
    ids = {r: upload(svc, [p])[0] for r, p in j.items()}
    run = svc.create_run("JOURNAL_UPLOAD_CHECK", None, "2026-09-30", ids, {}, "이담당", "preparer")
    for cid in ("A01", "T02", "J01", "A04", "F05"):
        assert check(run, cid)["status"] == "PASS", cid
    jb = samples["journal_bad"]
    ids = {r: upload(svc, [p])[0] for r, p in jb.items()}
    bad = svc.create_run("JOURNAL_UPLOAD_CHECK", None, "2026-09-30", ids, {}, "이담당", "preparer",
                         rerun_reason="오류 파일 테스트")
    assert check(bad, "A01")["status"] == "FAIL" and check(bad, "A01")["severity"] == "CRITICAL"
    assert check(bad, "J01")["status"] == "FAIL"
    assert check(bad, "A04")["status"] == "FAIL"
    assert bad["status"] == "BLOCKED"


def test_new_workflow_discovery_to_template_to_run(svc, samples):
    ids = upload(svc, samples["etf"].values())
    a = svc.analyze(ids)
    assert "plan" not in a
    nw = a["new_workflow"]
    assert nw["ok"] and nw["key_fields"] and nw["compare_fields"]
    assert any("확정하지 않습니다" in u for u in nw["uncertainties"])
    body = nw["template"]
    t = treg.save_version(svc.db, body, "이담당", "preparer", "ETF 잔고 대사 신규 등록")
    treg.set_status(svc.db, body["template_key"], t["version"], "ACTIVE", "박승인", "approver", "검토 후 활성화")
    a2 = svc.analyze(ids)
    cand = next(c for c in a2["candidates"] if c["template_key"] == body["template_key"])
    assert cand["is_candidate"]
    files = {r: v["file_id"] for r, v in cand["assignments"].items()}
    run = svc.create_run(body["template_key"], None, "2026-09-30", files, {}, "이담당", "preparer")
    assert check(run, "R01")["status"] == "FAIL"     # 예탁기관 누락
    assert check(run, "R02")["status"] == "FAIL"     # 예탁기관에만 존재
    assert check(run, "R05")["status"] == "FAIL"     # 수량 차이
    assert run["summary"]["different"] == 1


def test_pdf_prices_marked_for_review(svc, samples):
    p = dict(samples["current"])
    p["prices"] = samples["pdf"]
    run = bond_run(svc, p)
    assert check(run, "F08")["status"] == "REVIEW_REQUIRED"
    assert check(run, "R01")["status"] == "PASS"
    assert run["summary"]["processed_count"] == 22


def test_materiality_settings(svc):
    m = svc.set_materiality({"abs_amount": "5000"}, "관리자", "admin")
    assert m["abs_amount"] == "5000"
    with pytest.raises(ServiceError):
        svc.set_materiality({"pct": "-1"}, "관리자", "admin")
